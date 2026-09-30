import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


CONTACTS_FILE = Path("contacts.xlsx")
PROFILE_DIR = Path("whatsapp_profile")
SCREENSHOTS_DIR = Path("screenshots")
RECORDINGS_DIR = Path("recordings")
REPORT_DATE = date.today().isoformat()
DEFAULT_MESSAGE_TEMPLATE = "Hello {name}"
ACTION_PAUSE_MS = 3000
LOGIN_TIMEOUT_MS = 180_000


def create_contacts_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Contacts"
    worksheet.append(["Name", "Phone", "Message"])
    workbook.save(CONTACTS_FILE)


def load_contacts():
    if not CONTACTS_FILE.exists():
        create_contacts_template()
        print(f"Created {CONTACTS_FILE}. Add contacts, then run the script again.")
        return []

    workbook = load_workbook(CONTACTS_FILE, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError("contacts.xlsx must have a header row with Name and Phone.")

    header_indexes = {
        str(value).strip().casefold(): index
        for index, value in enumerate(headers)
        if value is not None
    }
    if "name" not in header_indexes or "phone" not in header_indexes:
        raise ValueError("contacts.xlsx must include Name and Phone columns.")

    contacts = []
    for row_number, row in enumerate(rows, start=2):
        name = row[header_indexes["name"]] if header_indexes["name"] < len(row) else None
        phone = row[header_indexes["phone"]] if header_indexes["phone"] < len(row) else None
        message = (
            row[header_indexes["message"]]
            if "message" in header_indexes and header_indexes["message"] < len(row)
            else None
        )
        if name is None and phone is None:
            continue

        name = str(name).strip() if name is not None else ""
        phone = str(phone).strip() if phone is not None else ""
        message = str(message).strip() if message is not None else ""
        if not name or not phone:
            contacts.append({
                "name": name or f"Row {row_number}",
                "phone": phone,
                "message_template": message,
                "input_error": "Name and Phone are both required.",
            })
            continue

        contacts.append({
            "name": name,
            "phone": phone,
            "message_template": message or DEFAULT_MESSAGE_TEMPLATE,
        })

    workbook.close()
    return contacts


def normalized_phone(phone):
    return re.sub(r"[\s()-]", "", phone)


def safe_filename(value):
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")
    return cleaned or "contact"


def find_visible_locator(page, selectors, timeout_ms=5000):
    for selector in selectors:
        locator = page.locator(selector).first
        try:
            locator.wait_for(state="visible", timeout=timeout_ms)
            return locator
        except PlaywrightTimeoutError:
            continue
    raise PlaywrightTimeoutError(f"Could not find a visible element matching: {selectors}")


def search_and_open_contact(page, contact):
    search_box = find_visible_locator(
        page,
        [
            'div[contenteditable="true"][data-tab="3"]',
            'div[contenteditable="true"][aria-label*="Search"]',
        ],
        timeout_ms=10_000,
    )
    for query in (normalized_phone(contact["phone"]), contact["name"]):
        search_box.fill(query)
        page.wait_for_timeout(ACTION_PAUSE_MS)
        result = page.get_by_text(query, exact=False).first
        try:
            result.wait_for(state="visible", timeout=5000)
            result.click()
            page.wait_for_timeout(ACTION_PAUSE_MS)
            return
        except PlaywrightTimeoutError:
            continue

    raise LookupError(f"Contact not found by phone or name: {contact['name']}")


def extract_last_incoming_messages(page):
    incoming = page.locator("div.message-in .selectable-text")
    messages = []
    for index in range(incoming.count()):
        text = incoming.nth(index).inner_text().strip()
        if text:
            messages.append(text)
    return messages[-3:]


def process_contact(page, contact):
    result = {
        "name": contact["name"],
        "phone": contact["phone"],
        "status": "failed",
        "message": "",
        "sent_at": None,
        "screenshot": None,
        "last_3_incoming_messages": [],
        "error": contact.get("input_error"),
    }
    if result["error"]:
        return result

    phone = normalized_phone(contact["phone"])
    if not re.fullmatch(r"\+\d{8,15}", phone):
        result["error"] = "Phone must include a country code, for example +919876543210."
        return result

    message = contact["message_template"].replace("{name}", contact["name"])
    result["message"] = message

    try:
        search_and_open_contact(page, contact)
        message_box = find_visible_locator(
            page,
            [
                'footer div[contenteditable="true"][data-tab="10"]',
                'footer [contenteditable="true"][aria-label="Type a message"]',
                'footer div[contenteditable="true"]',
            ],
            timeout_ms=10_000,
        )
        message_box.fill(message)
        message_box.press("Enter")

        outgoing_message = page.locator("div.message-out").filter(has_text=message).last
        outgoing_message.wait_for(state="visible", timeout=15_000)
        page.wait_for_timeout(ACTION_PAUSE_MS)

        SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)
        screenshot_path = SCREENSHOTS_DIR / f"{safe_filename(contact['name'])}_{safe_filename(phone)}.png"
        page.screenshot(path=str(screenshot_path))

        result["status"] = "sent"
        result["sent_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
        result["screenshot"] = str(screenshot_path)
        result["last_3_incoming_messages"] = extract_last_incoming_messages(page)
    except LookupError as error:
        result["error"] = str(error)
    except PlaywrightTimeoutError as error:
        result["error"] = f"Timed out waiting for WhatsApp Web: {error}"
    except Exception as error:
        result["error"] = f"{type(error).__name__}: {error}"

    return result


def save_reports(results):
    json_path = Path(f"whatsapp_report_{REPORT_DATE}.json")
    excel_path = Path(f"whatsapp_report_{REPORT_DATE}.xlsx")
    report = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "contacts_file": str(CONTACTS_FILE),
        "results": results,
    }
    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "WhatsApp Report"
    worksheet.append([
        "Name",
        "Phone",
        "Status",
        "Message",
        "Last 3 Incoming Messages",
        "Screenshot",
        "Error",
    ])
    for result in results:
        worksheet.append([
            result["name"],
            result["phone"],
            result["status"],
            result["message"],
            "\n".join(result["last_3_incoming_messages"]),
            result["screenshot"] or "",
            result["error"] or "",
        ])
    workbook.save(excel_path)
    print(f"Reports saved: {json_path} and {excel_path}")


def main():
    try:
        contacts = load_contacts()
    except Exception as error:
        print(f"Could not read contacts.xlsx: {error}")
        return

    if not contacts:
        return

    print(f"Loaded {len(contacts)} contact row(s) from {CONTACTS_FILE}.")
    print("Only proceed for people who have agreed to receive these messages.")
    if input("Open WhatsApp Web and prepare to send? Type 'yes' to continue: ").strip().casefold() != "yes":
        print("Cancelled. No messages were sent.")
        return

    results = []
    RECORDINGS_DIR.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            headless=False,
            record_video_dir=str(RECORDINGS_DIR),
            record_video_size={"width": 1280, "height": 720},
        )
        page = context.pages[0] if context.pages else context.new_page()
        page.goto("https://web.whatsapp.com", wait_until="domcontentloaded", timeout=60_000)
        print("If needed, scan the QR code. Waiting for the WhatsApp chat search field...")

        try:
            find_visible_locator(
                page,
                [
                    'div[contenteditable="true"][data-tab="3"]',
                    'div[contenteditable="true"][aria-label*="Search"]',
                ],
                timeout_ms=LOGIN_TIMEOUT_MS,
            )
            for contact in contacts:
                result = process_contact(page, contact)
                results.append(result)
                print(f"{result['name']}: {result['status']}")
                page.wait_for_timeout(ACTION_PAUSE_MS)
        except PlaywrightTimeoutError as error:
            print(f"WhatsApp Web did not become ready: {error}")
        finally:
            context.close()

    save_reports(results)


if __name__ == "__main__":
    main()